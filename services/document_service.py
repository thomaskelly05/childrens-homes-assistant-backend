from docx import Document
from io import BytesIO


def replace_placeholders(doc, data):
    for p in doc.paragraphs:
        for key, value in data.items():
            if f"{{{{{key}}}}}" in p.text:
                p.text = p.text.replace(f"{{{{{key}}}}}", str(value))


def generate_doc(template_path: str, data: dict):

    doc = Document(template_path)

    replace_placeholders(doc, data)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer
